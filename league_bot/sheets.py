"""Read-only client for the public Lowlands League spreadsheet."""

from __future__ import annotations

import csv
import io
import json
import re
import time
import urllib.parse
import urllib.request
from dataclasses import dataclass
from typing import Iterable

import openpyxl

DEFAULT_SPREADSHEET_ID = "15lCX6ljpsGh6bwzHZcctb24DIn-r4wqRGt6gZtLE_k4"


class SheetError(RuntimeError):
    """Raised when public sheet data cannot be loaded or interpreted."""


def normalized(value: object) -> str:
    return str(value or "").strip().casefold()


def number(value: object) -> float:
    raw = str(value or "").strip().replace(" ", "")
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
    elif "," in raw:
        decimals = len(raw.rsplit(",", 1)[1])
        raw = raw.replace(",", ".") if decimals != 3 else raw.replace(",", "")
    try:
        return float(raw)
    except ValueError:
        return 0.0


def display_number(value: object) -> str:
    parsed = number(value)
    return f"{parsed:,.0f}" if parsed.is_integer() else f"{parsed:,.1f}"


def first_value(row: dict[str, str], names: Iterable[str], default: str = "") -> str:
    lookup = {normalized(key): value for key, value in row.items()}
    for name in names:
        if normalized(name) in lookup:
            return str(lookup[normalized(name)] or "").strip()
    return default


@dataclass
class PublicSheet:
    spreadsheet_id: str = DEFAULT_SPREADSHEET_ID
    cache_seconds: int = 120

    def __post_init__(self) -> None:
        self._cache: dict[str, tuple[float, list[dict[str, str]]]] = {}
        self._profile_cache: tuple[float, dict[str, str]] | None = None
        self._user_cache: dict[str, tuple[float, dict[str, str]]] = {}
        self._tab_cache: tuple[float, dict[str, str]] | None = None

    def _request(self, url: str) -> bytes:
        request = urllib.request.Request(url, headers={"User-Agent": "Lowlands-League-Bot/2.0"})
        with urllib.request.urlopen(request, timeout=30) as response:
            return response.read()

    def profile_urls(self) -> dict[str, str]:
        if self._profile_cache and time.monotonic() - self._profile_cache[0] < 600:
            return self._profile_cache[1]
        url = f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/export?format=xlsx"
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(self._request(url)), read_only=False, data_only=False)
            sheet = workbook["Standings"]
            profiles = {}
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, 3)
                if cell.value and cell.hyperlink and cell.hyperlink.target:
                    profiles[normalized(cell.value)] = cell.hyperlink.target
        except Exception:
            profiles = {}
        self._profile_cache = (time.monotonic(), profiles)
        return profiles

    def profile(self, player: str) -> dict[str, str]:
        url = self.profile_urls().get(normalized(player), "")
        if not url:
            return {"url": "", "avatar": "", "country_code": ""}
        match = re.search(r"/user/([^/?#]+)", url)
        if not match:
            return {"url": url, "avatar": "", "country_code": ""}
        user_id = urllib.parse.unquote(match.group(1))
        cached = self._user_cache.get(user_id)
        if cached and time.monotonic() - cached[0] < 900:
            return cached[1]
        result = {"url": url, "avatar": "", "country_code": ""}
        try:
            data = json.loads(self._request(f"https://www.geoguessr.com/api/v3/users/{user_id}").decode())
            image = str(
                data.get("fullBodyPin")
                or (data.get("avatar") or {}).get("fullBodyPath")
                or (data.get("pin") or {}).get("url")
                or data.get("customImage")
                or ""
            ).lstrip("/")
            if image:
                result["avatar"] = (
                    "https://www.geoguessr.com/images/resize:fill:512:512/gravity:no:0:0/plain/" + image
                )
            result["country_code"] = str(data.get("countryCode") or "").lower()
        except Exception:
            pass
        self._user_cache[user_id] = (time.monotonic(), result)
        return result

    def tab_gids(self) -> dict[str, str]:
        if self._tab_cache and time.monotonic() - self._tab_cache[0] < 3600:
            return self._tab_cache[1]
        url = f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/edit"
        gids: dict[str, str] = {}
        try:
            page = self._request(url).decode("utf-8", errors="ignore")
            pattern = re.compile(
                r'\[21350203,"\[\d+,0,\\"(\d+)\\",\[\{\\"1\\":\[\[0,0,\\"([^\"]+)'
            )
            gids = {name.rstrip("\\"): gid for gid, name in pattern.findall(page)}
        except Exception:
            pass
        self._tab_cache = (time.monotonic(), gids)
        return gids

    def tab_url(self, tab: str) -> str:
        base = f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/edit"
        gid = self.tab_gids().get(tab)
        return f"{base}?gid={gid}#gid={gid}" if gid else base

    def rows(self, tab: str) -> list[dict[str, str]]:
        cached = self._cache.get(tab)
        if cached and time.monotonic() - cached[0] < self.cache_seconds:
            return cached[1]
        query = urllib.parse.urlencode({"tqx": "out:csv", "sheet": tab})
        url = f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/gviz/tq?{query}"
        request = urllib.request.Request(url, headers={"User-Agent": "Lowlands-League-Bot/1.0"})
        try:
            with urllib.request.urlopen(request, timeout=20) as response:
                text = response.read().decode("utf-8-sig")
        except Exception as exc:
            raise SheetError(f"Could not load the {tab} tab: {exc}") from exc
        rows = [dict(row) for row in csv.DictReader(io.StringIO(text))]
        if not rows:
            raise SheetError(f"The {tab} tab returned no rows.")
        self._cache[tab] = (time.monotonic(), rows)
        return rows

    def standings(self) -> list[dict[str, str]]:
        rows = []
        for source in self.rows("Standings"):
            player = first_value(source, ("Player", "Name"))
            rank = first_value(source, ("Rank", "#"))
            if not player or not rank:
                continue
            rows.append({
                "rank": rank,
                "player": player,
                "country": first_value(source, ("Country",)),
                "points": first_value(source, ("League Points", "Points", "Total"), "0"),
                "weeks": first_value(source, ("Weeks Played", "Weeks"), "0"),
                "moving": first_value(source, ("Moving",), "0"),
                "nm": first_value(source, ("NM",), "0"),
                "nmpz": first_value(source, ("NMPZ",), "0"),
                "playoffs": first_value(source, ("Playoffs",)),
                "profile_url": "",
            })
        profiles = self.profile_urls()
        for row in rows:
            row["profile_url"] = profiles.get(normalized(row["player"]), "")
        return rows

    def player(self, query: str) -> dict[str, str] | None:
        needle = normalized(query)
        exact = [row for row in self.standings() if normalized(row["player"]) == needle]
        if exact:
            return exact[0]
        partial = [row for row in self.standings() if needle in normalized(row["player"])]
        return partial[0] if len(partial) == 1 else None

    def all_player_stats(self) -> list[dict[str, object]]:
        rows = []
        for source in self.rows("Player Stats"):
            player = first_value(source, ("Player", "Name"))
            if not player:
                continue
            rows.append({
                "rank": first_value(source, ("Player Stats Rank", "Rank")),
                "player": player,
                "country": first_value(source, ("Country",)),
                "status": first_value(source, ("Status",)),
                "points": first_value(source, ("League Points",)),
                "moving": first_value(source, ("Moving",)),
                "nm": first_value(source, ("NM",)),
                "nmpz": first_value(source, ("NMPZ",)),
                "weeks": first_value(source, ("Weeks",)),
                "average": first_value(source, ("Average / Week",)),
                "last_three": first_value(source, ("Last 3",)),
                "consistency": first_value(source, ("Consistency",)),
                "best_week": first_value(source, ("Best Week",)),
                "worst_week": first_value(source, ("Worst Week",)),
                "best_mode": first_value(source, ("Best Mode",)),
                "cutoff_gap": first_value(source, ("Cutoff Gap",)),
                "streak": first_value(source, ("Top-16 Streak",)),
                "fives": first_value(source, ("5Ks", "Total 5Ks", "Five Ks")),
                "weekly": [first_value(source, (f"Week {week}",)) for week in range(1, 11)],
                "profile_url": "",
                "avatar": "",
            })
        profiles = self.profile_urls()
        for row in rows:
            row["profile_url"] = profiles.get(normalized(row["player"]), "")
        return rows

    def player_stats(self, query: str) -> dict[str, object] | None:
        rows = self.all_player_stats()
        needle = normalized(query)
        exact = [row for row in rows if normalized(row["player"]) == needle]
        if exact:
            return exact[0]
        partial = [row for row in rows if needle in normalized(row["player"])]
        return partial[0] if len(partial) == 1 else None

    def mode_leaderboard(self, mode: str) -> list[dict[str, str]]:
        key = normalized(mode)
        if key not in {"moving", "nm", "nmpz"}:
            return []
        return sorted(self.standings(), key=lambda row: number(row[key]), reverse=True)

    def season_insights(self) -> dict[str, dict[str, str]]:
        result: dict[str, dict[str, str]] = {}
        query = urllib.parse.urlencode({"tqx": "out:csv", "sheet": "Insights"})
        url = f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/gviz/tq?{query}"
        text = self._request(url).decode("utf-8-sig")
        for source in csv.reader(io.StringIO(text)):
            values = [str(value or "").strip() for value in source]
            for index, label in enumerate(values):
                key = normalized(label)
                if key in {"most 5ks", "moving specialist", "nm specialist", "nmpz specialist"}:
                    following = [value for value in values[index + 1 :] if value]
                    result[key] = {
                        "player": following[0] if following else "",
                        "value": following[1] if len(following) > 1 else "",
                    }
        return result

    def week(self, week: int) -> list[dict[str, str]]:
        result = []
        for source in self.rows("Challenges"):
            if int(number(first_value(source, ("Week",)))) != week:
                continue
            result.append({
                "number": first_value(source, ("Challenge #", "Challenge")),
                "mode": first_value(source, ("Mode",)),
                "map": first_value(source, ("Map",)),
                "deadline": first_value(source, ("Deadline",)),
                "status": first_value(source, ("Status",)),
                "url": first_value(source, ("Challenge URL", "URL", "Link")),
            })
        return result
